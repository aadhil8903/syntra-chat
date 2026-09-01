import pytest
import tempfile
import os
from docx import Document
from pypdf import PdfWriter
from rag.ingestion import extract_text_from_file
from core.gridfs_storage import get_gridfs_temp_file, fetch_gridfs_bytes

def test_extract_text_from_txt_file():
    with tempfile.NamedTemporaryFile(mode="w", suffix=".txt", delete=False, encoding="utf-8") as f:
        f.write("This is a sample document for testing RAG ingestion.\nSecond line.")
        temp_path = f.name

    try:
        results = extract_text_from_file(temp_path, "txt")
        assert len(results) >= 1
        assert "sample document" in results[0][0]
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

def test_extract_text_from_docx_file():
    doc = Document()
    doc.add_heading("Employee Benefits", level=1)
    doc.add_paragraph("Employees are eligible for health insurance starting on day one.")
    
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        temp_path = f.name
        doc.save(temp_path)

    try:
        results = extract_text_from_file(temp_path, "docx")
        assert len(results) >= 1
        full_text = " ".join([r[0] for r in results])
        assert "health insurance" in full_text
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

def test_extract_text_from_csv_file():
    with tempfile.NamedTemporaryFile(mode="w", suffix=".csv", delete=False, encoding="utf-8") as f:
        f.write("Name,Department,Salary\nAlice,Engineering,100000\nBob,HR,80000")
        temp_path = f.name

    try:
        results = extract_text_from_file(temp_path, "csv")
        assert len(results) >= 1
        assert "Alice" in results[0][0]
        assert "Engineering" in results[0][0]
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)

